
import os
import re

import pytest
from openpyxl import load_workbook
from playwright.sync_api import (
    Page,
    expect,
    TimeoutError as PlaywrightTimeoutError,
)

from pages.loginPage import LoginPage
from pages.profilesPage import ProfilesPage


@pytest.mark.ui
class TestProfilesPage:

    @pytest.fixture(autouse=True)
    def setup_profiles(self, page: Page, base_url: str):
        self.page = page

        email = os.getenv("TEST_EMAIL")
        password = os.getenv("TEST_PASSWORD")

        assert email, "Please set TEST_EMAIL."
        assert password, "Please set TEST_PASSWORD."

        page.set_default_navigation_timeout(60000)

        self.login = LoginPage(page)

        try:
            self.login.open(base_url)
        except PlaywrightTimeoutError:
            expect(
                page.get_by_placeholder("Enter email address")
            ).to_be_visible(timeout=15000)

        self.login.enter_email(email)
        self.login.enter_password(password)
        self.login.click_login()

        expect(page).to_have_url(
            re.compile(r".*/nolimit/app/home/?$"),
            timeout=45000
        )

        self.profiles = ProfilesPage(page)
        self.profiles.open_profiles()

    # ==========================================
    # 1. Profiles heading
    # ==========================================

    def test_profiles_heading(self):
        self.profiles.verify_heading()

    # ==========================================
    # 2. Description
    # ==========================================

    def test_profiles_description(self):
        self.profiles.verify_description()

    # ==========================================
    # 3. Search input
    # ==========================================

    def test_search_input_visible(self):
        self.profiles.verify_search_input()

    # ==========================================
    # 4. Search icon
    # ==========================================

    def test_search_icon_visible(self):
        self.profiles.verify_search_icon()

    # ==========================================
    # 5. Search placeholder
    # ==========================================

    def test_search_placeholder(self):
        self.profiles.verify_search_placeholder()

    # ==========================================
    # 6. Search accepts characters
    # ==========================================

    def test_search_accepts_characters(self):
        text = "Test123@"

        self.profiles.enter_search_text(text)

        expect(
            self.profiles.search_input
        ).to_have_value(text)

    # ==========================================
    # 7. Search by Client ID
    # ==========================================

    def test_search_by_client_id(self):
        client_id = self.profiles.get_first_row_value(0)

        print(f"Searching Client ID: {client_id}")

        self.profiles.perform_search(client_id)
        self.profiles.verify_search_results(
            0, client_id
        )

    # ==========================================
    # 8. Search by Name
    # ==========================================

    def test_search_by_name(self):
        name = self.profiles.get_first_row_value(1)

        print(f"Searching Name: {name}")

        self.profiles.perform_search(name)
        self.profiles.verify_search_results(
            1, name
        )

    # ==========================================
    # 9. Date picker visible
    # ==========================================

    def test_date_picker_visible(self):
        self.profiles.verify_date_picker()

    # ==========================================
    # 10. Date placeholders
    # ==========================================

    def test_date_placeholders(self):
        self.profiles.verify_date_placeholders()

    # ==========================================
    # 11. Calendar icon
    # ==========================================

    def test_calendar_icon_visible(self):
        self.profiles.verify_calendar_icon()

    # ==========================================
    # 12. Date selection
    # ==========================================

    def test_select_start_and_end_dates(self):
        start, end = self.profiles.select_date_range()

        print(f"Start Date: {start}")
        print(f"End Date: {end}")

        assert start, "Start Date is empty"
        assert end, "End Date is empty"

    # ==========================================
    # 13. Extract button
    # ==========================================

    def test_extract_button_visible(self):
        self.profiles.verify_extract_button()

    # ==========================================
    # 14. Extract icon
    # ==========================================

    def test_extract_down_arrow_icon(self):
        self.profiles.verify_extract_icon()

    # ==========================================
    # 15. Download Excel
    # ==========================================

    def test_extract_downloads_excel(self, tmp_path):
        self.profiles.select_my_files()

        saved_path = (
            self.profiles.download_profiles_excel(
                tmp_path
            )
        )

        print(f"Downloaded Excel: {saved_path.name}")

        assert saved_path.suffix.lower() == ".xlsx"

    # ==========================================
    # 16. Excel contains My Files records
    # ==========================================

    def test_excel_contains_my_files_records(self, tmp_path):
        self.profiles.select_my_files()

        table_records = self.profiles.get_all_records()

        assert table_records, (
            "My Files table contains no records"
        )

        saved_path = (
            self.profiles.download_profiles_excel(
                tmp_path
            )
        )

        workbook = load_workbook(
            saved_path,
            read_only=True,
            data_only=True
        )

        try:
            exported_values = set()

            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(
                    values_only=True
                ):
                    for value in row:
                        if value is not None:
                            exported_values.add(
                                str(value).strip()
                            )

            assert exported_values, (
                "Downloaded workbook is empty"
            )

            for record in table_records:
                client_id = record["client_id"]

                assert client_id in exported_values, (
                    f"Client ID {client_id} "
                    "is missing from Excel"
                )

        finally:
            workbook.close()

    # ==========================================
    # 17. Client ID header
    # ==========================================

    def test_client_id_header(self):
        self.profiles.verify_header(
            1, "Client ID"
        )

    # ==========================================
    # 18. Client ID value
    # ==========================================

    def test_client_id_value(self):
        client_id = self.profiles.get_first_row_value(0)

        assert re.fullmatch(
            r"SENTINEL[A-Z0-9]+",
            client_id
        ), f"Invalid Client ID: {client_id}"

    # ==========================================
    # 19. Client ID navigation
    # ==========================================

    def test_client_id_redirects_to_screening(self):
        client_id = (
            self.profiles.click_first_client_id()
        )

        self.profiles.verify_screening_navigation(
            client_id
        )

        print(f"Clicked Client ID: {client_id}")
        print(f"Redirected URL: {self.page.url}")

    # ==========================================
    # 20. Name header
    # ==========================================

    def test_name_header(self):
        self.profiles.verify_header(2, "Name")

    # ==========================================
    # 21. Name value
    # ==========================================

    def test_name_value(self):
        name = self.profiles.get_first_row_value(1)

        assert name.strip(), "Name is empty"

    # ==========================================
    # 22. Type header
    # ==========================================

    def test_type_header(self):
        self.profiles.verify_header(3, "Type")

    # ==========================================
    # 23. Type value
    # ==========================================

    def test_type_value(self):
        record_type = self.profiles.get_first_row_value(2)

        assert record_type.casefold() in {
            "individual",
            "entity"
        }, f"Unexpected Type: {record_type}"

    # ==========================================
    # 24. Status header
    # ==========================================

    def test_status_header(self):
        self.profiles.verify_header(4, "Status")

    # ==========================================
    # 25. Status value
    # ==========================================

    def test_status_value(self):
        status = self.profiles.get_first_row_value(3)

        assert status.casefold() == "onboarded", (
            f"Unexpected Status: {status}"
        )

    # ==========================================
    # 26. Onboarded Date header
    # ==========================================

    def test_onboarded_date_header(self):
        self.profiles.verify_header(
            5, "Onboarded Date"
        )

    # ==========================================
    # 27. Onboarded Date value
    # ==========================================

    def test_onboarded_date_value(self):
        value = self.profiles.get_first_row_value(4)

        assert re.fullmatch(
            r"\d{2}-\d{2}-\d{4}",
            value
        ), f"Invalid Onboarded Date: {value}"

    # ==========================================
    # 28. Risk Rating header
    # ==========================================

    def test_risk_rating_header(self):
        self.profiles.verify_header(
            6, "Risk Rating"
        )

    # ==========================================
    # 29. Risk Rating value
    # ==========================================

    def test_risk_rating_value(self):
        rating = self.profiles.get_first_row_value(5)

        assert rating.casefold() in {
            "low",
            "medium",
            "high"
        }, f"Unexpected Risk Rating: {rating}"

    # ==========================================
    # 30. Onboarded By header
    # ==========================================

    def test_onboarded_by_header(self):
        self.profiles.verify_header(
            7, "Onboarded By"
        )

    # ==========================================
    # 31. Onboarded By value
    # ==========================================

    def test_onboarded_by_value(self):
        value = self.profiles.get_first_row_value(6)

        assert value.strip(), (
            "Onboarded By is empty"
        )

    # ==========================================
    # 32. My Files shows logged-in user's records
    # ==========================================

    def test_my_files_shows_logged_in_user_records(self):
        self.profiles.select_my_files()

        records = self.profiles.get_all_records()

        assert records, (
            "No records found under My Files"
        )

        # Temporary account-specific expected name.
        # Replace with a dynamic account-name locator
        # once the header DOM is inspected.
        logged_in_user = "Oshadha Weerakon"

        for record in records:
            actual_user = record["onboarded_by"]

            print(
                f"Onboarded By: {actual_user}"
            )

            assert (
                actual_user.casefold()
                == logged_in_user.casefold()
            ), (
                f"My Files contains a record onboarded "
                f"by '{actual_user}', expected "
                f"'{logged_in_user}'"
            )

    # ==========================================
    # 33. My Files folder icon
    # ==========================================

    def test_my_files_folder_icon(self):
        self.profiles.select_my_files()
        self.profiles.verify_my_files_icon()

    # ==========================================
    # 34. All Files shows My Files records too
    # ==========================================

    def test_all_files_shows_all_users_records(self):
        self.profiles.select_my_files()

        my_records = self.profiles.get_all_records()

        self.profiles.select_all_files()

        all_records = self.profiles.get_all_records()

        my_ids = {
            record["client_id"]
            for record in my_records
        }

        all_ids = {
            record["client_id"]
            for record in all_records
        }

        assert my_ids.issubset(all_ids), (
            "Some My Files records are missing "
            "from All Files"
        )

        print(f"My Files count: {len(my_ids)}")
        print(f"All Files count: {len(all_ids)}")

    # ==========================================
    # 35. All Files folder icon
    # ==========================================

    def test_all_files_folder_icon(self):
        self.profiles.select_all_files()
        self.profiles.verify_all_files_icon()
